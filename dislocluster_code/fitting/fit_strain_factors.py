"""
fit_strain_factors.py — calibrate the growth strain factors A_a, A_c to measured
irradiation-growth strain.

Model:  eps_a =  A_a*(CiL_i + CaiL_i)        (a-elongation, interstitial a-loops)
        eps_c = -A_c*(CvL_v + CavL_v)        (c-contraction, vacancy c-loops)
so A_a, A_c are linear scalings of the (already calibrated) loop content. With the
loop content vs dose fixed by the model, each factor is a least-squares scaling
through the origin:  A = sum(content*|eps_target|) / sum(content^2).

TARGETS — representative annealed alpha-Zr irradiation growth (order-of-magnitude,
~550-600 K), from the review literature in docs/ (Holt 1988 'Mechanisms of
Irradiation Growth'; Griffiths neutron-damage reviews; Carpenter/Northwood). Zr
growth is ~volume-conserving (eps_a+eps_b+eps_c=0, basal-isotropic => eps_c=-2*eps_a),
with a low-dose transient and c-loop 'breakaway' at higher dose. Replace with a
specific dataset (material/texture/T) when available.
"""
import io, contextlib
from datetime import datetime
import numpy as np
from dislocluster_code import paths as _paths
from dislocluster_code.zerod.simulation import ZrMicroSimulation
from dislocluster_code.zerod.cpp_bridge import run_cpp_solver
_N = contextlib.redirect_stdout(io.StringIO())

# (dose [dpa], eps_a target).  eps_c target = -2*eps_a (volume conservation).
STRAIN_TARGETS = [
    (1.0,  1.0e-3),
    (5.0,  3.0e-3),
    (10.0, 6.0e-3),
]
T_FIT = 573.0   # K (workbook default; representative growth-data temperature)

with _N:
    # was the cwd-relative 'input/Zr_input_parameters.xlsx', which only
    # resolved when the script happened to be launched from ZrMicro/
    SIM = ZrMicroSimulation(_paths.INPUT_DIR / 'Zr_input_parameters.xlsx')
inp = SIM.input_data
inp.material_params['T'] = T_FIT
with _N:
    inp.calculate_derived_parameters()
G = inp.material_params['G']
doses = np.array([d for d, _ in STRAIN_TARGETS])
cfg = {'t_begin': 1e-1, 't_end': max(doses)/G*1.5, 'n_points': 400,
       'rtol': 1e-6, 'atol': 1e-20, 'log_time': True,
       'solver_method': {'backend': 'cvode', 'lmm': 'bdf', 'linsol': 'dense'}}
with _N:
    r = run_cpp_solver(SIM, cfg, base_dir=_paths.ZRMICRO_DIR)
c = r['concentrations']
dose = G * r['time']

# loop content (atom fraction) at the target doses
def interp(dp, arr):
    return float(np.interp(dp, dose, arr))
content_a = np.array([interp(d, c['CiL_i'] + c['CaiL_i']) for d, _ in STRAIN_TARGETS])
content_v = np.array([interp(d, c['CvL_v'] + c['CavL_v']) for d, _ in STRAIN_TARGETS])
eps_a_t = np.array([e for _, e in STRAIN_TARGETS])
eps_c_t = 2.0 * eps_a_t   # |eps_c| = 2*eps_a

# least-squares scaling through origin
A_a = float(np.sum(content_a * eps_a_t) / np.sum(content_a**2))
A_c = float(np.sum(content_v * eps_c_t) / np.sum(content_v**2))

print("=" * 64)
print(f"STRAIN-FACTOR CALIBRATION (annealed Zr anchors, T={T_FIT:.0f} K)")
print("=" * 64)
print(f"  A_a = {A_a:.4g}   (was 1.0)")
print(f"  A_c = {A_c:.4g}   (was 0.5)")
print(f"\n{'dpa':>6} | {'content_a':>10} {'eps_a mdl':>10} {'eps_a tgt':>10} | "
      f"{'content_v':>10} {'eps_c mdl':>10} {'eps_c tgt':>10}")
for i, (d, _) in enumerate(STRAIN_TARGETS):
    print(f"{d:>6.1f} | {content_a[i]:>10.3e} {A_a*content_a[i]:>10.3e} {eps_a_t[i]:>10.3e} | "
          f"{content_v[i]:>10.3e} {-A_c*content_v[i]:>10.3e} {-eps_c_t[i]:>10.3e}")
print(f"\nA_c/A_a = {A_c/A_a:.1f}  (large ratio reflects model i-loop content >> "
      f"v-loop content;\n  growth volume-conservation eps_c=-2eps_a is NOT enforced by the model)")

# write into the workbook
import openpyxl
_WB = _paths.INPUT_DIR / 'Zr_input_parameters.xlsx'
wb = openpyxl.load_workbook(_WB)
ws = wb['Model_Parameters']
def setp(key, val):
    for row in range(1, ws.max_row+1):
        if str(ws.cell(row,1).value).strip() == key:
            old = ws.cell(row,4).value; ws.cell(row,4).value = val
            print(f"  Model_Parameters {key}: {old} -> {val:.4g}"); return
    raise KeyError(key)
print("\nApplying to workbook:")
setp('A_a', A_a); setp('A_c', A_c)
wb.save(_WB)
print("saved.")
